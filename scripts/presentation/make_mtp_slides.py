import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

output_path = "/root/Code-Base/SSS/MTP_Farm_ERP_Presentation.pptx"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

C_PRIMARY = RGBColor(15, 76, 129)       # Classic Navy
C_SECONDARY = RGBColor(30, 130, 76)     # Emerald Green
C_ACCENT = RGBColor(230, 126, 34)       # Golden Amber
C_DARK = RGBColor(33, 37, 41)           # Charcoal Text
C_WHITE = RGBColor(255, 255, 255)
C_CARD_BORDER = RGBColor(222, 226, 230)
C_CARD_BG = RGBColor(255, 255, 255)

blank_slide_layout = prs.slide_layouts[6]

def add_header(slide, title_text, category="MTP FARM ERP • HE THONG QUAN TRI TOAN DIEN"):
    header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
    tf = header_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p_cat = tf.paragraphs[0]
    p_cat.text = category.upper()
    p_cat.font.size = Pt(10)
    p_cat.font.bold = True
    p_cat.font.color.rgb = C_SECONDARY
    
    p_title = tf.add_paragraph()
    p_title.text = title_text
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = C_PRIMARY
    p_title.space_before = Pt(4)

def add_card(slide, left, top, width, height, title, body_bullets, tag=None, tag_color=C_PRIMARY, bg_color=C_CARD_BG):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.color.rgb = C_CARD_BORDER
    shape.line.width = Pt(1)

    tb = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), width - Inches(0.5), height - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    if tag:
        p_tag = tf.paragraphs[0]
        p_tag.text = tag.upper()
        p_tag.font.size = Pt(9)
        p_tag.font.bold = True
        p_tag.font.color.rgb = tag_color
        
        p_title = tf.add_paragraph()
        p_title.text = title
        p_title.font.size = Pt(13)
        p_title.font.bold = True
        p_title.font.color.rgb = C_PRIMARY
        p_title.space_before = Pt(2)
    else:
        p_title = tf.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(13)
        p_title.font.bold = True
        p_title.font.color.rgb = C_PRIMARY

    for bullet in body_bullets:
        p = tf.add_paragraph()
        p.text = "• " + bullet
        p.font.size = Pt(10)
        p.font.color.rgb = C_DARK
        p.space_before = Pt(4)

# SLIDE 1: COVER
s1 = prs.slides.add_slide(blank_slide_layout)
bg = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
bg.fill.solid()
bg.fill.fore_color.rgb = RGBColor(12, 35, 64)
bg.line.fill.background()

tb = s1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(4.0))
tf = tb.text_frame
tf.word_wrap = True

p1 = tf.paragraphs[0]
p1.text = "🐔 MTP FARM ERP"
p1.font.size = Pt(38)
p1.font.bold = True
p1.font.color.rgb = RGBColor(245, 171, 53)

p2 = tf.add_paragraph()
p2.text = "Hệ Thống Quản Trị Toàn Diện Trại Gà Đẻ Trứng Minh Tân Phát"
p2.font.size = Pt(22)
p2.font.bold = True
p2.font.color.rgb = C_WHITE
p2.space_before = Pt(8)

p3 = tf.add_paragraph()
p3.text = "Đơn vị áp dụng: Trang Trại Gà Đẻ Trứng Lê Văn Dương (Minh Tân, Dầu Tiếng, Bình Dương)"
p3.font.size = Pt(12.5)
p3.font.color.rgb = RGBColor(200, 215, 230)
p3.space_before = Pt(18)

p4 = tf.add_paragraph()
p4.text = "Số hóa Vật tư • Quản lý Cấp phát • Trạm Dầu Xe Cơ Giới • Chống Thất Thoát • Báo Cáo Chi Phí Chuồng Trại"
p4.font.size = Pt(11.5)
p4.font.color.rgb = RGBColor(180, 200, 220)
p4.space_before = Pt(6)

# SLIDE 2: BỐI CẢNH & VẤN ĐỀ
s2 = prs.slides.add_slide(blank_slide_layout)
add_header(s2, "1. Bối Cảnh & Bài Toán Quản Trị Trang Trại Quy Mô Lớn")

add_card(s2, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0), 
         "Thất Thoát Vật Tư Cơ Điện", 
         ["Hàng trăm loại động cơ quạt hút, motor bạt, bóng đèn sưởi thay thế liên tục.",
          "Khó kiểm soát xác thiết bị cũ hỏng, dễ bị tuồn bán phế liệu ra ngoài.",
          "Không có quy trình thu hồi đổi 1-1 minh bạch giữa thợ và kho."],
         "VẤN ĐỀ #1", RGBColor(220, 53, 69))

add_card(s2, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0), 
         "Gian Lận & Thất Thoát Dầu", 
         ["Trại sử dụng hàng chục xe cơ giới (xe tải, xe nâng, cày) & trạm máy phát dự phòng.",
          "Việc đổ dầu vào xe hoặc máy phát điện khó đo lường chính xác mức tiêu hao.",
          "Nguy cơ cấp khống số lít, hao hụt bồn chứa không rõ nguyên nhân."],
         "VẤN ĐỀ #2", RGBColor(220, 53, 69))

add_card(s2, Inches(8.8), Inches(1.8), Inches(3.6), Inches(5.0), 
         "Chi Phí Mù Mờ Theo Chuồng", 
         ["Không biết chính xác chuồng A1, A2 hay B1 ngốn bao nhiêu tiền vật tư hàng tháng.",
          "Ghi chép sổ sách giấy tờ dễ thất lạc, số liệu chốt cuối tháng bị sai lệch.",
          "Kế toán mất nhiều ngày đối soát chứng từ xuất kho."],
         "VẤN ĐỀ #3", RGBColor(220, 53, 69))

# SLIDE 3: GIẢI PHÁP TỔNG THỂ
s3 = prs.slides.add_slide(blank_slide_layout)
add_header(s3, "2. Giải Pháp MTP Farm ERP: Single Source of Truth")

add_card(s3, Inches(0.8), Inches(1.8), Inches(5.6), Inches(2.4),
         "Đồng Bộ Thời Gian Thực (Realtime)",
         ["Mọi giao dịch xuất nhập kho, cấp dầu, đổi đồ hỏng cập nhật ngay lập tức.",
          "Loại bỏ hoàn toàn sổ sách giấy tờ và nhập liệu trung gian."],
         "TÍNH LIỀN MẠCH", C_SECONDARY)

add_card(s3, Inches(6.8), Inches(1.8), Inches(5.6), Inches(2.4),
         "Phân Cấp & Xác Nhận 2 Chiều",
         ["Trưởng chuồng yêu cầu -> Quản lý/Thủ kho duyệt -> Người nhận bấm xác nhận đã nhận đủ.",
          "Trách nhiệm rõ ràng, không thể đổ lỗi thất lạc đồ."],
         "MINH BẠCH", C_PRIMARY)

add_card(s3, Inches(0.8), Inches(4.5), Inches(5.6), Inches(2.4),
         "Báo Cáo Chi Phí Chuồng Trại Chi Tiết",
         ["Mỗi que hàn, bóng đèn, lít dầu đều gán trực tiếp vào mã dãy chuồng phụ trách.",
          "Ban Giám đốc nắm rõ tỷ suất sinh lời và chi phí vận hành từng chuồng."],
         "QUẢN TRỊ CHI PHÍ", C_ACCENT)

add_card(s3, Inches(6.8), Inches(4.5), Inches(5.6), Inches(2.4),
         "PWA & Hỗ Trợ Offline Sóng Yếu",
         ["Cài đặt như app Native trên điện thoại của quản lý và công nhân.",
          "Giao diện tối ưu thao tác nhanh ngay tại hiện trường chuồng trại."],
         "TRẢI NGHIỆM TIỆN LỢI", C_SECONDARY)

# SLIDE 4: KIẾN TRÚC KỸ THUẬT
s4 = prs.slides.add_slide(blank_slide_layout)
add_header(s4, "3. Kiến Trúc Kỹ Thuật & Công Nghệ Nền Tảng")

add_card(s4, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0),
         "Frontend & Giao Diện",
         ["Framework: Next.js 15 (App Router) siêu tốc (~200ms).",
          "Ngôn ngữ: TypeScript Strict Mode an toàn dữ liệu.",
          "Styling: Tailwind CSS v4 & Lucide Icons hiện đại.",
          "Mobile PWA: Hỗ trợ Add-to-Homescreen trên iOS/Android."],
         "FRONTEND LAYER", C_PRIMARY)

add_card(s4, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0),
         "Backend & Database",
         ["Cơ sở dữ liệu: Supabase PostgreSQL 17 mạnh mẽ.",
          "Bảo mật: Row Level Security (RLS) phân quyền dữ liệu.",
          "Realtime: WebSocket push thông báo chuông tức thì.",
          "Authentication: Username Auth tối ưu cho công nhân trại."],
         "BACKEND & DATA", C_SECONDARY)

add_card(s4, Inches(8.8), Inches(1.8), Inches(3.6), Inches(5.0),
         "Chất Lượng & Triển Khai",
         ["Automated Tests: Hơn 318 test cases pass 100% (Vitest).",
          "Containerization: Docker & Docker Compose sẵn sàng.",
          "CI/CD: Tự động hóa kiểm thử và đóng gói bản phát hành.",
          "Backup: Cơ chế tự động backup CSDL định kỳ."],
         "DEVOPS & TESTING", C_ACCENT)

# SLIDE 5: VẬT TƯ & ĐỔI ĐỒ 1-1
s5 = prs.slides.add_slide(blank_slide_layout)
add_header(s5, "4. Phân Hệ Quản Lý Vật Tư, Kho Bãi & Cơ Chế Đổi Đồ 1-1")

add_card(s5, Inches(0.8), Inches(1.8), Inches(5.6), Inches(2.4),
         "Danh Mục Đa Biến Thể & Bộ Composite",
         ["Quản lý vật tư theo biến thể (công suất, điện áp 220V/380V, hãng).",
          "Composite Kits: 1 bộ quạt hút xuất trọn gói hoặc xuất lẻ từng linh kiện."],
         "DANH MỤC VẬT TƯ", C_PRIMARY)

add_card(s5, Inches(6.8), Inches(1.8), Inches(5.6), Inches(2.4),
         "Cơ Chế Đổi Đồ Cũ - Lấy Đồ Mới 1-1",
         ["Bắt buộc thu hồi xác motor, bóng đèn, bo mạch hỏng khi cấp đồ mới.",
          "Quản lý kho xác phế liệu riêng, thanh lý định kỳ có hạch toán thu hồi vốn."],
         "CHỐNG THẤT THOÁT", C_SECONDARY)

add_card(s5, Inches(0.8), Inches(4.5), Inches(5.6), Inches(2.4),
         "Mượn Trả Dụng Cụ Đồ Nghề Cơ Điện",
         ["Quản lý máy hàn, máy khoan, đồng hồ đo điện cấp cho thợ.",
          "Hệ thống tự động gửi thông báo chuông nhắc nhở đồ mượn quá hạn."],
         "QUẢN LÝ TÀI SẢN", C_ACCENT)

add_card(s5, Inches(6.8), Inches(4.5), Inches(5.6), Inches(2.4),
         "Cảnh Báo Tồn Kho Tối Thiểu (Min/Max)",
         ["Tự động đổi màu cảnh báo vàng/đỏ khi vật tư quan trọng sắp hết.",
          "Đảm bảo trại không bị gián đoạn điện lưới hay quạt làm mát."],
         "AN TOÀN VẬN HÀNH", C_SECONDARY)

# SLIDE 6: TRẠM DẦU & XE CƠ GIỚI
s6 = prs.slides.add_slide(blank_slide_layout)
add_header(s6, "5. Phân Hệ Quản Lý Trạm Cấp Dầu & Xe Cơ Giới")

add_card(s6, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0),
         "Nhập Bồn Dầu Tổng",
         ["Ghi nhận số lít nhập, nhà cung cấp, hóa đơn và đơn giá.",
          "Theo dõi biến động giá dầu và tính toán tồn kho bồn chính xác theo thời gian thực."],
         "QUẢN LÝ BỒN CHỨA", C_PRIMARY)

add_card(s6, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0),
         "Cấp Phát Theo Xe & Máy Phát",
         ["Định danh từng xe cơ giới (xe tải, xe ben, xe nâng) và trạm máy phát điện.",
          "Ghi nhận số giờ chạy máy (hour-meter) hoặc số km trước khi đổ."],
         "ĐỊNH MỨC CƠ GIỚI", C_SECONDARY)

add_card(s6, Inches(8.8), Inches(1.8), Inches(3.6), Inches(5.0),
         "Cảnh Báo Tiêu Hao Bất Thường",
         ["Tự động tính định mức lít/giờ hoặc lít/km.",
          "Báo động ngay cho Ban Quản Lý khi có xe đổ dầu với mức tiêu hao vượt ngưỡng cho phép."],
         "KIỂM SOÁT THẤT THOÁT", RGBColor(220, 53, 69))

# SLIDE 7: BÁO CÁO CHI PHÍ THEO CHUỒNG
s7 = prs.slides.add_slide(blank_slide_layout)
add_header(s7, "6. Báo Cáo Doanh Thu, Chi Phí & Hạch Toán Dãy Chuồng")

add_card(s7, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0),
         "Phân Bổ Chi Phí Từng Chuồng Gà",
         ["Hạch toán chi tiết chi phí vật tư cơ điện cho từng chuồng (A1, A2, B1...).",
          "So sánh định mức tiêu hao vật tư giữa các dãy chuồng cùng quy mô.",
          "Phát hiện chuồng có thiết bị hư hỏng bất thường để bảo trì phòng ngừa.",
          "Xuất báo cáo tài chính chi tiết theo định dạng Excel/PDF."],
         "COST ACCOUNTING", C_PRIMARY)

add_card(s7, Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0),
         "Dashboard Tổng Quan Dành Cho Chủ Trại",
         ["Biểu đồ phân bổ chi phí vật tư: Điện, Nước, Cơ điện, Nhiên liệu.",
          "Theo dõi tỷ lệ hoàn thành phiếu yêu cầu vật tư đúng hạn.",
          "Thống kê tổng giá trị tài sản thu hồi từ phế liệu và sửa chữa ngoài.",
          "Truy cập an toàn trên mọi thiết bị máy tính, iPad và smartphone."],
         "EXECUTIVE DASHBOARD", C_SECONDARY)

# SLIDE 8: LỢI ÍCH & TỔNG KẾT
s8 = prs.slides.add_slide(blank_slide_layout)
add_header(s8, "7. Giá Trị Thực Tế & Lợi Ích Cho Doanh Nghiệp")

add_card(s8, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0),
         "Tiết Kiệm 15 - 25% Chi Phí",
         ["Chặn đứng thất thoát vật tư và gian lận xăng dầu xe cơ giới.",
          "Tận dụng tối đa giá trị thu hồi từ phế liệu và xác thiết bị cũ."],
         "TỐI ƯU CHI PHÍ", C_SECONDARY)

add_card(s8, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0),
         "Nâng Cao 50% Năng Suất",
         ["Thủ kho và trưởng chuồng duyệt xuất vật tư trong 30 giây.",
          "Không tốn thời gian ghi chép giấy tờ hay đối soát thủ công cuối tháng."],
         "TỐC ĐỘ VẬN HÀNH", C_PRIMARY)

add_card(s8, Inches(8.8), Inches(1.8), Inches(3.6), Inches(5.0),
         "Chuyển Đổi Số Toàn Diện",
         ["Xây dựng nền tảng quản trị nông nghiệp công nghệ cao hiện đại.",
          "Dễ dàng mở rộng thêm các phân hệ: Sản lượng trứng, Thức ăn cám, Thú y."],
         "TƯƠNG LAI BỀN VỮNG", C_ACCENT)

# SLIDE 9: CLOSING
s9 = prs.slides.add_slide(blank_slide_layout)
bg9 = s9.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
bg9.fill.solid()
bg9.fill.fore_color.rgb = RGBColor(12, 35, 64)
bg9.line.fill.background()

tb9 = s9.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.3), Inches(3.5))
tf9 = tb9.text_frame
tf9.word_wrap = True

p_end1 = tf9.paragraphs[0]
p_end1.text = "🐔 MTP FARM ERP"
p_end1.font.size = Pt(36)
p_end1.font.bold = True
p_end1.font.color.rgb = RGBColor(245, 171, 53)
p_end1.alignment = PP_ALIGN.CENTER

p_end2 = tf9.add_paragraph()
p_end2.text = "Số Hóa Vận Hành • Nâng Tầm Nông Nghiệp Việt"
p_end2.font.size = Pt(20)
p_end2.font.color.rgb = C_WHITE
p_end2.alignment = PP_ALIGN.CENTER
p_end2.space_before = Pt(10)

p_end3 = tf9.add_paragraph()
p_end3.text = "Trang Trại Gà Đẻ Trứng Lê Văn Dương • Minh Tân, Dầu Tiếng, Bình Dương"
p_end3.font.size = Pt(13)
p_end3.font.color.rgb = RGBColor(180, 200, 220)
p_end3.alignment = PP_ALIGN.CENTER
p_end3.space_before = Pt(25)

os.makedirs(os.path.dirname(output_path), exist_ok=True)
prs.save(output_path)
print("SUCCESS: Presentation created at " + output_path)
